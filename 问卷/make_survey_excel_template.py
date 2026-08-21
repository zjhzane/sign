# -*- coding: utf-8 -*-
"""生成问卷批量填写用的 Excel 模板（与 auto_fill_survey.py 列名一致）。"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


def desktop_dir() -> Path:
    local = Path.home() / "Desktop"
    if local.exists():
        return local.resolve()
    cn = Path.home() / "桌面"
    if cn.exists():
        return cn.resolve()
    return (Path.home() / "Desktop").resolve()


def main() -> None:
    out = desktop_dir() / "问卷批量填写_学号模板.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "学号名单"

    headers = ["学号", "姓名", "毕业去向", "用人单位名称", "是否QS100以内"]
    head_fill = PatternFill("solid", fgColor="D5DDE5")
    head_font = Font(name="微软雅黑", size=11, bold=True, color="1C2B36")
    body_font = Font(name="微软雅黑", size=11, color="1C2B36")
    thin = Border(
        left=Side(style="thin", color="A8B4C0"),
        right=Side(style="thin", color="A8B4C0"),
        top=Side(style="thin", color="A8B4C0"),
        bottom=Side(style="thin", color="A8B4C0"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = thin

    samples = [
        ("2022010001", "张三", "境外留学", "曼彻斯特大学", "是"),
        ("2022010002", "李四", "境外留学", "麦考瑞大学", "否"),
    ]
    for r, row in enumerate(samples, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = body_font
            cell.alignment = center if c != 4 else left
            cell.border = thin

    for r in range(4, 34):
        for c in range(1, 6):
            cell = ws.cell(r, c, None)
            cell.font = body_font
            cell.alignment = center if c != 4 else left
            cell.border = thin

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 18
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:E33"

    dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    dv.error = "请填写 是 或 否"
    dv.errorTitle = "是否QS100以内"
    dv.prompt = "填 是 或 否（对应问卷重点高校/QS前100）"
    dv.promptTitle = "QS100"
    ws.add_data_validation(dv)
    dv.add("E2:E200")

    ws["A1"].comment = Comment(
        "脚本必读：学号、是否QS100以内。\n姓名/毕业去向/用人单位名称仅便于核对。\n是否QS100以内填：是 / 否。",
        "template",
    )

    note = wb.create_sheet("填写说明")
    note["A1"] = "问卷批量填写 Excel 格式说明"
    note["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    lines = [
        "",
        "1. 在「学号名单」从第 2 行填写。第 2–3 行是示例，用前请删掉或改成真人。",
        "2. 列含义：",
        "   学号：必填。脚本按这一列逐人登录问卷。",
        "   姓名：选填，只用于日志核对。",
        "   毕业去向：选填。当前脚本默认就业状态=境外留学。",
        "   用人单位名称：选填，院校名称，便于人工对照 QS。",
        "   是否QS100以内：填 是 或 否。对应问卷「我升学的院校是否属于重点高校（双高/双一流或QS前100）或重点科研机构」。",
        "3. 运行：python 问卷/auto_fill_survey.py --excel 本文件路径",
        "4. 脚本默认也会读桌面「新建 Microsoft Excel 工作表.xlsx」，列名保持一致即可。",
    ]
    for i, line in enumerate(lines, 2):
        note.cell(i, 1, line).font = Font(name="微软雅黑", size=11)
    note.column_dimensions["A"].width = 110

    wb.save(out)
    print(out.resolve())


if __name__ == "__main__":
    main()
