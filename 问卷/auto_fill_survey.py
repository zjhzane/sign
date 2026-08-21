# -*- coding: utf-8 -*-
"""
浙江万里学院 2026届毕业生就业质量与人才培养调查问卷
Playwright 自动填写脚本

默认：
  - 从 Excel 读学号批量填写（填完一个再下一个）
  - 毕业去向「境外留学」→ 第1题选境外留学
  - 毕业去向「签劳动合同 / 签就业协议形式就业」→ 第1题选境内全职工作；地点/规模/行业/属于按表填写；单位性质默认民营企业
  - 毕业去向「升学—研究生」→ 第1题选境内升学；重点高校/QS前100 按表格 E 列填写
  - 毕业去向「自主创业—创立公司 / 个体工商户」→ 第1题选自主创业；「创业类型」按破折号后的类型填写
  - 「是否重点高校/QS前100」按表格「是否QS100以内」列填 是/否
  - 多选题随机 1–3 项（不选「其他」）
  - 满意度类选最高档（很满意 / 5分 / 很符合 等）
  - 注意力检测题选「比较不符合」
  - 默认不自动提交（加 --submit 才会提交）

用法:
  pip install playwright openpyxl
  playwright install chromium
  python auto_fill_survey.py
  python auto_fill_survey.py --excel "C:/Users/zjhza/Desktop/问卷批量填写_学号模板.xlsx" --submit
  python auto_fill_survey.py --student-id 2022015409 --submit
  python auto_fill_survey.py --headless
"""

from __future__ import annotations

import argparse
import random
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, sync_playwright

URL = "https://dm.njcedu.com/answer?id=17093"
DEFAULT_STUDENT_ID = "2022015557"
DEFAULT_EXCEL = r"C:\Users\zjhza\Desktop\问卷批量填写_学号模板.xlsx"
EMPLOYMENT_STATUS = "境外留学"
STATUS_FULLTIME = "境内全职工作"
ATTENTION_OPTION = "比较不符合"
DEFAULT_NATURE = "民营企业"
# 职业级联无表格列时，按物流/商务类岗位优先匹配
DEFAULT_JOB_KEYS = "物流 运输 仓储 邮政 商务 销售 商业和服务业 办事人员"


@dataclass
class StudentRow:
    student_id: str
    name: str = ""
    destination: str = ""  # 毕业去向
    unit_name: str = ""
    qs100: str = "否"
    address: str = ""
    industry: str = ""
    belong: str = ""  # 属于：小微企业等
    scale: str = ""
    nature: str = DEFAULT_NATURE


def log(msg: str) -> None:
    print(msg, flush=True)
    path = getattr(log, "_path", None)
    if path:
        with open(path, "a", encoding="utf-8") as f:
            f.write(msg + "\n")


def get_title(page: Page) -> str:
    loc = page.locator(".p-title")
    if loc.count():
        t = loc.first.inner_text().strip()
        t = re.sub(
            r"^(单选题|多选题|填空题|组合题|级联题|填空|编码题|检索题)\s*",
            "",
            t,
        ).strip()
        return t
    return ""


def get_type_badge(page: Page) -> str:
    loc = page.locator(".p-wtType")
    if loc.count():
        return loc.first.inner_text().strip()
    # fallback from title area
    loc = page.locator(".p-title")
    if loc.count():
        t = loc.first.inner_text().strip()
        m = re.match(r"^(单选题|多选题|填空题|组合题)", t)
        if m:
            return m.group(1)
    return ""


def radio_labels(page: Page):
    return page.locator("label.el-radio:visible")


def checkbox_labels(page: Page):
    return page.locator("label.el-checkbox:visible")


def option_texts(locator) -> List[str]:
    texts = []
    for i in range(locator.count()):
        try:
            t = locator.nth(i).inner_text().strip()
            t = re.sub(r"\s+", " ", t)
            if t:
                texts.append(t)
        except Exception:
            pass
    return texts


def click_label_with_text(locator, text: str) -> bool:
    for i in range(locator.count()):
        el = locator.nth(i)
        try:
            t = re.sub(r"\s+", " ", el.inner_text().strip())
            if t == text or text in t:
                el.click(timeout=3000)
                return True
        except Exception:
            continue
    return False


def is_attention(title: str) -> bool:
    return "比较不符合" in title and ("不要选其他" in title or "本题请选择" in title)


def is_negative_item(title: str) -> bool:
    hints = [
        "很困难",
        "不喜欢",
        "忘记答应",
        "马马虎虎",
        "比较强势",
        "独自行动",
        "不太考虑",
        "转移话题",
        "讨厌",
        "避免和自己意见不同",
        "不适合高强度",
        "消极影响",
        "只注意自己感兴趣",
        "希望有人",
        "不愿意推荐" if False else "",  # placeholder
    ]
    return any(h and h in title for h in hints)


def normalize_yes_no(val) -> str:
    s = str(val or "").strip()
    if s in ("是", "Y", "y", "1", "true", "True", "YES", "yes"):
        return "是"
    return "否"


def is_key_uni_question(title: str) -> bool:
    return "重点高校" in title or "QS前" in title or "重点科研" in title


def _cell(row, idx: Optional[int]) -> str:
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def is_contract_job(student: Optional[StudentRow]) -> bool:
    if not student:
        return False
    d = student.destination or ""
    return (
        "劳动合同" in d
        or "签劳动合同" in d
        or "就业协议" in d
        or "签就业协议" in d
    )


def is_domestic_study(student: Optional[StudentRow]) -> bool:
    if not student:
        return False
    d = student.destination or ""
    if "境外" in d or "出国" in d:
        return False
    return "境内升学" in d or "升学" in d or "研究生" in d


def is_startup(student: Optional[StudentRow]) -> bool:
    if not student:
        return False
    return "自主创业" in (student.destination or "")


def venture_type_of(student: Optional[StudentRow]) -> str:
    """从毕业去向解析创业类型：创立公司 / 个体工商户。"""
    if not student:
        return ""
    d = (student.destination or "").replace("—", "-").replace("－", "-").replace("–", "-")
    if "个体工商户" in d:
        return "个体工商户"
    if "创立公司" in d or "创办公司" in d:
        return "创立公司"
    return ""


def employment_status_of(student: Optional[StudentRow]) -> str:
    if not student:
        return EMPLOYMENT_STATUS
    if is_contract_job(student):
        return STATUS_FULLTIME
    d = student.destination or ""
    if "境外留学" in d:
        return "境外留学"
    if is_domestic_study(student):
        return "境内升学"
    if "自由职业" in d:
        return "自由职业"
    if is_startup(student):
        return "自主创业"
    if "未就业" in d:
        return "未就业"
    return EMPLOYMENT_STATUS


def norm_text(s: str) -> str:
    t = (s or "").strip()
    t = t.replace("一下", "以下").replace(" ", "")
    t = t.replace("（", "(").replace("）", ")")
    return t


def split_address(addr: str) -> List[str]:
    s = (addr or "").replace(" ", "")
    if not s:
        return []
    parts: List[str] = []
    rest = s
    for pat in (
        r".+?(?:自治区|特别行政区|省)",
        r".+?市",
        r".+?(?:区|县|旗)",
    ):
        while rest:
            m = re.match(pat, rest)
            if not m:
                break
            parts.append(m.group(0))
            rest = rest[m.end() :]
            # 省只切一次，市通常一次，区可能一次
            if "省" in pat or "自治区" in pat:
                break
            if pat == r".+?市":
                break
    if rest:
        parts.append(rest)
    tokens: List[str] = []
    for p in parts:
        tokens.append(p)
        short = re.sub(r"(省|自治区|特别行政区|市|州|盟|区|县|旗)$", "", p)
        if short and short != p:
            tokens.append(short)
    seen = set()
    out: List[str] = []
    for t in tokens:
        if t not in seen:
            seen.add(t)
            out.append(t)
    return out


def best_option_match(options: List[str], want: str) -> Optional[str]:
    if not want or not options:
        return None
    w = norm_text(want)
    for o in options:
        if norm_text(o) == w:
            return o
    # 选项包含完整目标（如「交通运输、仓储和邮政业」）
    contain = [o for o in options if w in norm_text(o)]
    if contain:
        return max(contain, key=lambda o: len(norm_text(o)))

    pieces = [p.strip() for p in re.split(r"[、，,/及和]", w) if len(p.strip()) >= 2]
    for p in pieces:
        for o in options:
            no = norm_text(o)
            if p == no or p in no:
                return o

    # 规模：49人及以下 ↔ 50人以下
    nums = re.findall(r"\d+", w)
    if nums and ("人" in w or "规模" in w):
        n = int(nums[0])
        for o in options:
            no = norm_text(o)
            if str(n) in no:
                return o
            if n <= 50 and ("50人以下" in no or "50人及以下" in no or "49人" in no):
                return o
    if "小微" in w:
        for o in options:
            no = norm_text(o)
            if "小微" in no or "微型" in no or "小型" in no:
                return o
    return None


def visible_option_texts(page: Page) -> List[str]:
    loc = page.locator(
        ".el-cascader-node:visible .el-cascader-node__label, "
        ".el-select-dropdown__item:visible, "
        ".el-cascader-menu:visible .el-cascader-node:visible"
    )
    return option_texts(loc)


def click_visible_option(page: Page, text: str) -> bool:
    loc = page.locator(
        ".el-cascader-node:visible .el-cascader-node__label, "
        ".el-select-dropdown__item:visible, "
        ".el-cascader-menu:visible .el-cascader-node:visible"
    )
    return click_label_with_text(loc, text) or click_label_with_text(
        page.locator(".el-cascader-node:visible, .el-select-dropdown__item:visible"),
        text,
    )


def load_students_from_excel(path: str) -> List[StudentRow]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=1, max_row=500, max_col=16, values_only=True))
    wb.close()
    if not rows:
        return []

    header = [str(c or "").strip() for c in rows[0]]

    def col(*names: str) -> Optional[int]:
        for i, h in enumerate(header):
            for n in names:
                if n in h:
                    return i
        return None

    i_id = col("学号")
    i_name = col("姓名")
    i_dest = col("毕业去向", "去向")
    i_unit = col("用人单位名称", "单位名称")
    i_qs = col("是否QS100", "QS100", "重点高校")
    i_addr = col("用人单位地址", "工作地点", "地址")
    i_ind = col("行业")
    i_belong = col("属于")
    i_scale = col("规模")
    i_nature = col("单位性质", "性质")
    if i_id is None:
        raise ValueError(f"Excel 未找到「学号」列，表头={header}")

    out: List[StudentRow] = []
    for row in rows[1:]:
        if not row or i_id >= len(row) or row[i_id] is None:
            continue
        sid = str(row[i_id]).strip()
        if not sid or sid.lower() == "none":
            continue
        if re.fullmatch(r"\d+\.0", sid):
            sid = sid[:-2]
        dest = _cell(row, i_dest)
        nature = _cell(row, i_nature)
        if not nature:
            nature = "个体工商户" if "个体工商户" in dest else DEFAULT_NATURE
        out.append(
            StudentRow(
                student_id=sid,
                name=_cell(row, i_name),
                destination=dest,
                unit_name=_cell(row, i_unit),
                qs100=normalize_yes_no(_cell(row, i_qs)) if i_qs is not None else "否",
                address=_cell(row, i_addr),
                industry=_cell(row, i_ind),
                belong=_cell(row, i_belong),
                scale=_cell(row, i_scale),
                nature=nature,
            )
        )
    return out


def pick_radio(
    title: str,
    options: List[str],
    qs100: Optional[str] = None,
    student: Optional[StudentRow] = None,
) -> Optional[str]:
    if not options:
        return None

    if "就业状态" in title:
        want = employment_status_of(student)
        for o in options:
            if want in o:
                return o

    if student and ("创业类型" in title or "创业形式" in title):
        want = venture_type_of(student)
        if want:
            hit = best_option_match(options, want)
            if hit:
                return hit
            for o in options:
                if want in o:
                    return o

    if student and is_domestic_study(student) and (
        "升学类型" in title or "升学层次" in title or "升学形式" in title
    ):
        for pref in ("研究生", "硕士", "硕士研究生"):
            hit = best_option_match(options, pref)
            if hit:
                return hit
            for o in options:
                if pref in o:
                    return o

    if is_attention(title):
        for o in options:
            if ATTENTION_OPTION in o:
                return o

    # 重点高校 / QS前100：按表格填写
    if is_key_uni_question(title) and qs100 is not None:
        want = normalize_yes_no(qs100)
        for o in options:
            if o.strip() == want:
                return o
        for o in options:
            if want in o:
                return o

    if student:
        if "单位性质" in title or ("性质" in title and "单位" in title):
            hit = best_option_match(options, student.nature or DEFAULT_NATURE)
            if hit:
                return hit
            hit = best_option_match(options, "民营")
            if hit:
                return hit
        if "规模" in title:
            hit = best_option_match(options, student.scale)
            if hit:
                return hit
        if "属于" in title or "企业划型" in title or (
            "企业类型" in title and "创业" not in title
        ):
            hit = best_option_match(options, student.belong)
            if hit:
                return hit
        if "行业" in title:
            hit = best_option_match(options, student.industry)
            if hit:
                return hit

    # 高科技/重点产业：无表格列时选否（避免被后面的「是」优先级误伤）
    if "高科技" in title or "重点行业产业" in title:
        for o in options:
            if o.strip() == "否":
                return o

    # 离职
    if "离职" in title:
        if is_contract_job(student) or is_startup(student):
            for pref in ("0次", "从未离职", "0 次", "没有离职"):
                for o in options:
                    if pref in o:
                        return o
            for o in options:
                if "从未就业" not in o and ("0" in o or "无" in o):
                    return o
        else:
            for o in options:
                if "从未就业" in o:
                    return o

    # 健康消极影响
    if "健康" in title and "影响" in title:
        for o in options:
            if "没有影响" in o:
                return o

    # 反向表述 → 很不符合
    if is_negative_item(title):
        for pref in ("很不符合", "不符合", "没有影响"):
            for o in options:
                if o == pref:
                    return o

    # 最高满意度 / 最高分 / 最正向
    priority = [
        "很满意",
        "很吻合",
        "很胜任",
        "很符合",
        "非常符合",
        "愿意",
        "推荐报考本校本专业",
        "了解",
        "相关",
        "较多",
        "是",
        "5分",
        "有很大帮助",
        "显著加分，成为核心竞争优势",
        "满意",
        "吻合",
        "胜任",
        "符合",
        "比较了解",
        "基本相关",
        "有较大帮助",
    ]
    for pref in priority:
        for o in options:
            if "无法评价" in o:
                continue
            if o == pref or o.startswith(pref):
                return o

    # 社保类选保障最多
    for o in options:
        if "其他保障" in o or "企业年金" in o:
            return o

    for o in options:
        if "无法评价" not in o and o != "其他":
            return o
    return options[0]


def fill_radio(page: Page, title: str, student: Optional[StudentRow] = None) -> str:
    labs = radio_labels(page)
    opts = option_texts(labs)
    qs100 = student.qs100 if student else None
    best = pick_radio(title, opts, qs100=qs100, student=student)
    if best and click_label_with_text(labs, best):
        return best
    if opts and click_label_with_text(labs, opts[0]):
        return opts[0]
    return ""


def fill_multiple(page: Page, title: str) -> str:
    labs = checkbox_labels(page)
    opts = option_texts(labs)
    if not opts:
        return ""

    # 多选：排除「其他」「以上均未…」，随机选 1–3 项
    pool = [
        o
        for o in opts
        if not o.startswith("其他") and "以上均未" not in o and o.strip() != "其他"
    ]
    if not pool:
        return ""

    n = random.randint(1, min(3, len(pool)))
    # 题干限制「1-5项」时仍落在 1–3
    if "1-5" in title or "1－5" in title or "选出1" in title:
        n = random.randint(1, min(3, len(pool)))

    chosen = []
    for o in random.sample(pool, n):
        if click_label_with_text(labs, o):
            chosen.append(o)
            page.wait_for_timeout(100)
    return ",".join(chosen)


def fill_combine(page: Page) -> str:
    """组合矩阵表：每行点第一列（通常为很满意/5分）。"""
    clicked = 0
    tables = page.locator("table.cus-table:visible")
    if tables.count() == 0:
        tables = page.locator("table:visible")

    for t in range(tables.count()):
        table = tables.nth(t)
        rows = table.locator("tr")
        for r in range(rows.count()):
            row = rows.nth(r)
            # 跳过表头
            if row.locator("th").count():
                continue
            radios = row.locator("label.el-radio")
            if radios.count() == 0:
                continue
            try:
                radios.first.click(timeout=2000)
                clicked += 1
                page.wait_for_timeout(80)
            except Exception:
                try:
                    # 点内部 span
                    row.locator(".el-radio__inner").first.click(timeout=2000)
                    clicked += 1
                except Exception:
                    pass
    return f"rows={clicked}"


def fill_blank(page: Page, title: str) -> str:
    value = "10"
    if any(k in title for k in ("收入", "薪资", "期望")):
        value = "8000"
    elif "工作小时" in title or ("每周" in title and "工作" in title):
        value = "40"
    elif "课堂以外" in title or ("每周" in title and "学习" in title):
        value = str(random.randint(8, 15))

    inputs = page.locator(
        "input.el-input__inner:visible, textarea:visible, input[type='number']:visible, input:visible"
    )
    filled = []
    for i in range(inputs.count()):
        inp = inputs.nth(i)
        try:
            cls = inp.get_attribute("class") or ""
            typ = inp.get_attribute("type") or ""
            if "radio" in cls or typ == "radio" or typ == "checkbox":
                continue
            ph = inp.get_attribute("placeholder") or ""
            if "搜索" in ph:
                continue
            inp.fill(value)
            filled.append(value)
        except Exception:
            continue
    return ",".join(filled)


def _node_text(el) -> str:
    try:
        lab = el.locator(".el-cascader-node__label")
        if lab.count():
            t = lab.first.inner_text().strip()
        else:
            t = el.inner_text().strip()
        return re.sub(r"\s+", " ", t)
    except Exception:
        return ""


def cascader_last_menu(page: Page):
    menus = page.locator(".el-cascader-menu:visible")
    n = menus.count()
    if n == 0:
        return None
    return menus.nth(n - 1)


def menu_option_texts(menu) -> List[str]:
    loc = menu.locator(".el-cascader-node:visible")
    texts: List[str] = []
    for i in range(loc.count()):
        t = _node_text(loc.nth(i))
        if t:
            texts.append(t)
    return texts


def last_menu_option_texts(page: Page) -> List[str]:
    menu = cascader_last_menu(page)
    if menu is None:
        return visible_option_texts(page)
    return menu_option_texts(menu)


def click_menu_option(menu, text: str) -> bool:
    items = menu.locator(".el-cascader-node:visible")
    for i in range(items.count()):
        el = items.nth(i)
        t = _node_text(el)
        if not t:
            continue
        if t == text or text in t or (len(t) >= 2 and t in text):
            try:
                lab = el.locator(".el-cascader-node__label")
                if lab.count():
                    lab.first.click(timeout=3000)
                else:
                    el.click(timeout=3000)
                return True
            except Exception:
                continue
    return False


def click_dropdown_item(page: Page, text: str) -> bool:
    menu = cascader_last_menu(page)
    if menu is not None:
        return click_menu_option(menu, text)
    items = page.locator(".el-select-dropdown__item:visible, .el-cascader-node:visible")
    for i in range(items.count()):
        el = items.nth(i)
        t = _node_text(el)
        if not t:
            continue
        if t == text or text in t or (len(t) >= 2 and t in text):
            try:
                el.click(timeout=3000)
                return True
            except Exception:
                continue
    return False


def _want_keys(want: str, kind: str) -> List[str]:
    want = (want or "").strip()
    if not want:
        return []
    keys: List[str] = [want]
    if kind == "address" or re.search(r"(省|市|.+区|县|自治区)", want):
        for t in split_address(want):
            if t not in keys:
                keys.append(t)
    for t in re.split(r"[、，,/及和]", want):
        t = t.strip()
        if t and t not in keys:
            keys.append(t)
    return keys


def fill_code(page: Page, want: str = "", kind: str = "") -> str:
    """级联/下拉：按 want 从左到右匹配每一列；已选中的层级不再点击。"""
    steps: List[str] = []
    keys = _want_keys(want, kind)

    is_cascader = page.locator(".el-cascader:visible").count() > 0
    trigger = page.locator(".el-cascader:visible, .el-select:visible")
    if trigger.count() == 0:
        return ""
    try:
        trigger.first.click()
        page.wait_for_timeout(450)
    except Exception:
        return ""

    for level in range(8):
        if is_cascader:
            menus = page.locator(".el-cascader-menu:visible")
            if menus.count() <= level:
                break
            menu = menus.nth(level)
            opts = menu_option_texts(menu)
        else:
            menu = None
            opts = visible_option_texts(page)
        if not opts:
            break
        picked = None
        for k in keys:
            picked = best_option_match(opts, k)
            if picked:
                break
        if not picked:
            picked = opts[0]

        already = False
        if menu is not None:
            active = menu.locator(
                ".el-cascader-node.in-active-path, .el-cascader-node.is-active"
            )
            if active.count():
                at = _node_text(active.first)
                if at == picked or picked in at or at in picked:
                    already = True
        if already:
            steps.append(picked)
            continue

        ok = click_menu_option(menu, picked) if menu is not None else click_dropdown_item(page, picked)
        if not ok:
            break
        steps.append(picked)
        page.wait_for_timeout(400)
        if not is_cascader:
            break
        if page.locator(".el-cascader-menu:visible").count() == 0:
            break
        if page.locator(".el-cascader-menu:visible").count() <= level + 1:
            arrow = page.locator(
                ".el-cascader-menu:visible .el-cascader-node.in-active-path .el-cascader-node__postfix, "
                ".el-cascader-menu:visible .el-cascader-node.is-active .el-cascader-node__postfix"
            )
            if arrow.count():
                try:
                    arrow.last.click(timeout=2000)
                    page.wait_for_timeout(400)
                except Exception:
                    pass
            if page.locator(".el-cascader-menu:visible").count() <= level + 1:
                break

    if page.locator(".el-cascader-menu:visible").count():
        try:
            page.keyboard.press("Escape")
            page.wait_for_timeout(150)
        except Exception:
            pass
    return " > ".join(steps)


def fill_api(page: Page, name: str = "") -> str:
    name = (name or "").strip() or "宁波诺丁汉大学"
    inp = page.locator("input.el-input__inner:visible, input:visible").first
    try:
        inp.click()
        inp.fill(name)
        page.wait_for_timeout(700)
        sug = page.locator(
            ".el-autocomplete-suggestion li:visible, .el-select-dropdown__item:visible"
        )
        if sug.count():
            texts = option_texts(sug)
            hit = best_option_match(texts, name)
            if hit:
                click_label_with_text(sug, hit)
            else:
                page.keyboard.press("Escape")
        return name
    except Exception:
        return ""


def answer_one(page: Page, student: Optional[StudentRow] = None) -> str:
    badge = get_type_badge(page)
    title = get_title(page)
    short = title[:50].replace("\n", " ")
    qs100 = student.qs100 if student else None

    if is_attention(title) or ("比较不符合" in title and "不要选其他" in title):
        labs = radio_labels(page)
        if click_label_with_text(labs, ATTENTION_OPTION):
            return f"{short} => 注意力检测:{ATTENTION_OPTION}"

    if "就业状态" in title:
        want = employment_status_of(student)
        labs = radio_labels(page)
        if click_label_with_text(labs, want):
            return f"{short} => {want}"

    if student and ("创业类型" in title or "创业形式" in title):
        want = venture_type_of(student)
        labs = radio_labels(page)
        if want and click_label_with_text(labs, want):
            return f"{short} [单选] => {want}(表)"
        opts = option_texts(labs)
        hit = best_option_match(opts, want) if want else None
        if hit and click_label_with_text(labs, hit):
            return f"{short} [单选] => {hit}(表)"

    if student and is_domestic_study(student) and (
        "升学类型" in title or "升学层次" in title or "升学形式" in title
    ):
        labs = radio_labels(page)
        for pref in ("研究生", "硕士", "硕士研究生"):
            if click_label_with_text(labs, pref):
                return f"{short} [单选] => {pref}(表)"

    if is_key_uni_question(title) and qs100 is not None:
        want = normalize_yes_no(qs100)
        labs = radio_labels(page)
        if click_label_with_text(labs, want):
            return f"{short} [单选] => {want}(表)"

    if "组合" in badge:
        detail = fill_combine(page)
        return f"{short} [组合] => {detail}"

    if "多选" in badge or (checkbox_labels(page).count() > 0 and radio_labels(page).count() == 0):
        detail = fill_multiple(page, title)
        return f"{short} [多选] => {detail}"

    loc_q = any(
        k in title
        for k in (
            "所在地",
            "工作地点",
            "单位地点",
            "就业地点",
            "创业地点",
            "升学地点",
            "院校地点",
            "地址",
        )
    )
    industry_q = "行业" in title
    job_q = "职业" in title or "岗位类别" in title or "从事的职业" in title

    if "单位名称" in title or (
        "名称" in title
        and any(k in title for k in ("单位", "企业", "公司", "创业", "院校", "学校", "高校", "升学", "留学"))
    ):
        detail = fill_api(page, student.unit_name if student else "")
        return f"{short} [名称] => {detail}"

    if "填空" in badge:
        detail = fill_blank(page, title)
        return f"{short} [填空] => {detail}"

    # 下拉：性质 / 规模 / 属于 / 行业
    if page.locator(".el-select:visible").count() and radio_labels(page).count() == 0:
        want = ""
        kind = ""
        if student:
            if "创业类型" in title or "创业形式" in title:
                want = venture_type_of(student)
            elif "性质" in title:
                want = student.nature or DEFAULT_NATURE
            elif "规模" in title:
                want = student.scale
            elif "属于" in title:
                want = student.belong
            elif industry_q:
                want = student.industry
            elif loc_q:
                want = student.address
                kind = "address"
        detail = fill_code(page, want, kind=kind)
        return f"{short} [下拉] => {detail}"

    if page.locator(".el-cascader:visible").count() and radio_labels(page).count() == 0:
        want = ""
        kind = ""
        if student:
            if loc_q:
                want = student.address
                kind = "address"
            elif industry_q:
                want = student.industry
            elif job_q:
                want = DEFAULT_JOB_KEYS.replace(" ", ",")
        detail = fill_code(page, want, kind=kind)
        return f"{short} [级联] => {detail}"

    if radio_labels(page).count():
        detail = fill_radio(page, title, student=student)
        return f"{short} [单选] => {detail}"

    if checkbox_labels(page).count():
        detail = fill_multiple(page, title)
        return f"{short} [多选] => {detail}"

    if page.locator("input.el-input__inner:visible").count():
        if page.locator(".el-cascader:visible").count():
            want = ""
            kind = ""
            if student:
                if loc_q:
                    want = student.address
                    kind = "address"
                elif industry_q:
                    want = student.industry
                elif job_q:
                    want = DEFAULT_JOB_KEYS.replace(" ", ",")
            detail = fill_code(page, want, kind=kind)
            return f"{short} [级联] => {detail}"
        if "单位名称" in title or ("名称" in title and page.locator(".el-autocomplete:visible").count()):
            detail = fill_api(page, student.unit_name if student else "")
            return f"{short} [名称] => {detail}"
        detail = fill_blank(page, title)
        return f"{short} [填空] => {detail}"

    return f"{short} => 未能识别"


def dismiss_overlays(page: Page) -> None:
    """关掉级联/下拉遮罩，避免挡住登录和下一题。"""
    try:
        page.keyboard.press("Escape")
    except Exception:
        pass
    try:
        page.locator("text=拼命为您服务中").wait_for(state="hidden", timeout=2500)
    except Exception:
        pass


def page_toast(page: Page) -> str:
    texts: List[str] = []
    loc = page.locator(
        ".el-message__content, .el-message-box__message, .el-notification__content"
    )
    try:
        n = loc.count()
    except Exception:
        n = 0
    for i in range(n):
        try:
            t = loc.nth(i).inner_text().strip()
            t = re.sub(r"\s+", " ", t)
            if t and t not in texts:
                texts.append(t)
        except Exception:
            pass
    return "；".join(texts)


def click_next(page: Page, allow_submit: bool) -> str:
    page.wait_for_timeout(200)
    dismiss_overlays(page)

    # 滚到底，避免按钮被挡
    try:
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
    except Exception:
        pass
    page.wait_for_timeout(100)

    candidates = [
        page.locator("button.p-next-btn"),
        page.get_by_role("button", name="下一题"),
        page.get_by_role("button", name="下一页"),
        page.locator("button:has-text('下一题')"),
    ]
    for loc in candidates:
        try:
            if loc.count() and loc.first.is_visible():
                loc.first.scroll_into_view_if_needed()
                loc.first.click()
                try:
                    page.locator("text=拼命为您服务中").wait_for(state="visible", timeout=800)
                    page.locator("text=拼命为您服务中").wait_for(state="hidden", timeout=15000)
                except Exception:
                    page.wait_for_timeout(400)
                return "next"
        except Exception:
            continue

    if allow_submit:
        for name in ("提交", "确认提交", "完成"):
            sub = page.get_by_role("button", name=name)
            if sub.count() and sub.first.is_visible():
                sub.first.click()
                return "submit"
    # 再检查是否其实还有下一题但不可见
    any_next = page.locator("button.p-next-btn, button:has-text('下一题')")
    if any_next.count():
        try:
            any_next.first.click(force=True)
            page.wait_for_timeout(600)
            return "next"
        except Exception:
            pass
    return "none"


def _login_once(page: Page, student_id: str) -> str:
    """单次登录尝试。返回 pager / ended / unknown。"""
    page.goto(URL, wait_until="domcontentloaded", timeout=30000)
    page.wait_for_timeout(700)
    dismiss_overlays(page)

    if "surveyEnd" in page.url:
        return "ended"

    start = page.get_by_role("button", name="开始填写")
    try:
        if start.count() and start.first.is_visible():
            start.first.click()
            page.wait_for_timeout(700)
    except Exception:
        pass

    tab = page.get_by_role("tab", name="学号")
    try:
        if tab.count():
            tab.first.click()
            page.wait_for_timeout(250)
    except Exception:
        pass

    inp = page.get_by_placeholder("请输入学号")
    try:
        inp.first.wait_for(state="visible", timeout=8000)
    except PlaywrightTimeoutError:
        toast = page_toast(page)
        log(f"  登录表单未出现 url={page.url}" + (f" toast={toast}" if toast else ""))
        return "unknown"

    inp.first.click()
    inp.first.fill("")
    inp.first.fill(student_id)
    page.wait_for_timeout(200)

    confirm = page.get_by_role("button", name="确认")
    if confirm.count():
        confirm.first.click()
    else:
        page.keyboard.press("Enter")

    try:
        page.wait_for_url(re.compile(r"(pager|surveyEnd)"), timeout=25000)
    except PlaywrightTimeoutError:
        pass

    toast = page_toast(page)
    if "surveyEnd" in page.url:
        return "ended"
    if "pager" in page.url:
        page.wait_for_selector(
            "label.el-radio, label.el-checkbox, .p-title, button.p-next-btn",
            timeout=15000,
        )
        return "pager"
    if toast:
        log(f"  登录提示: {toast}")
        if any(k in toast for k in ("已填写", "已完成", "已经答过", "无需重复")):
            return "ended"
    return "unknown"


def login(page: Page, student_id: str, retries: int = 3) -> str:
    """登录并进入答题页。失败会重试，避免偶发未跳转。"""
    last = "unknown"
    for attempt in range(1, retries + 1):
        last = _login_once(page, student_id)
        if last in ("pager", "ended"):
            return last
        log(f"  登录未成功 ({attempt}/{retries}) url={page.url}")
        page.wait_for_timeout(800 * attempt)
        try:
            page.reload(wait_until="domcontentloaded")
        except Exception:
            pass
    return last


def wait_question_change(page: Page, old_title: str, timeout_ms: int = 8000) -> bool:
    end = time.time() + timeout_ms / 1000
    while time.time() < end:
        page.wait_for_timeout(200)
        body = page.locator("body").inner_text()
        if any(k in body for k in ("提交成功", "感谢您的参与", "答卷已提交", "提交完成")):
            return True
        new_title = get_title(page)
        if new_title and new_title != old_title:
            return True
        # 进度条变化也算
    return False


def fill_one_student(
    page: Page,
    student: StudentRow,
    submit: bool,
    max_steps: int,
) -> str:
    """填写一名学生，返回 ended / done / error。"""
    log(
        f"—— 开始：{student.student_id}"
        + (f" {student.name}" if student.name else "")
        + f"，去向={student.destination or '-'}，状态={employment_status_of(student)}"
        + f"，提交={submit}"
    )
    status = login(page, student.student_id)
    if status == "ended":
        log("该学号问卷已填写完成（surveyEnd），跳过。")
        return "ended"
    if status != "pager":
        log(f"登录后状态异常: url={page.url}")
        return "error"
    log(f"已进入答题页: {page.url}")

    for step in range(1, max_steps + 1):
        title_before = get_title(page)
        body = page.locator("body").inner_text()
        if any(k in body for k in ("提交成功", "感谢您的参与", "答卷已提交")):
            log("问卷已完成/已提交。")
            return "done"

        result = answer_one(page, student=student)
        log(f"[{step:03d}] {result}")
        page.wait_for_timeout(250)

        action = click_next(page, allow_submit=submit)
        if action == "none":
            page.wait_for_timeout(1200)
            dismiss_overlays(page)
            action = click_next(page, allow_submit=submit)
        if action == "submit":
            page.wait_for_timeout(2000)
            log("已点击提交。")
            return "done"
        if action == "none":
            if page.get_by_role("button", name=re.compile(r"提交")).count():
                log("已到末题。未加 --submit 则不提交。")
                return "done"
            log("找不到下一题按钮，停止该学号。")
            return "error"

        changed = wait_question_change(page, title_before, 10000)
        if not changed:
            log("  题目未跳转，重试作答…")
            answer_one(page, student=student)
            page.wait_for_timeout(300)
            click_next(page, allow_submit=submit)
            if not wait_question_change(page, title_before, 8000):
                log("仍无法跳转，停止该学号。")
                return "error"
    else:
        log(f"达到最大步数 {max_steps}")
        return "error"


def run(
    students: List[StudentRow],
    headless: bool,
    submit: bool,
    slow_mo: int,
    max_steps: int,
    keep_open_ms: int = 20000,
    log_file: str = "",
) -> None:
    if log_file:
        log._path = log_file  # type: ignore[attr-defined]
        open(log_file, "w", encoding="utf-8").close()
    else:
        log._path = None  # type: ignore[attr-defined]

    if not students:
        log("没有待填写的学号。")
        return

    def launch_browser(pw):
        return pw.chromium.launch(headless=headless, slow_mo=slow_mo)

    def open_page(br):
        ctx = br.new_context(viewport={"width": 1280, "height": 900})
        pg = ctx.new_page()
        pg.set_default_timeout(12000)
        return pg

    def close_page(pg) -> None:
        try:
            pg.context.close()
        except Exception:
            pass

    def page_alive(pg) -> bool:
        try:
            return not pg.is_closed()
        except Exception:
            return False

    def ensure_page(br, pg):
        if br is None or not br.is_connected():
            log("浏览器已关闭，正在重新打开…")
            try:
                if br is not None:
                    br.close()
            except Exception:
                pass
            br = launch_browser(p)
            return br, open_page(br)
        if not page_alive(pg):
            return br, open_page(br)
        return br, pg

    def fill_and_collect(pg, student, failed_list) -> None:
        try:
            result = fill_one_student(pg, student, submit=submit, max_steps=max_steps)
            if result == "error":
                failed_list.append(student)
        except PlaywrightTimeoutError as e:
            log(f"超时，跳过该学号: {e}")
            failed_list.append(student)
        except Exception as e:
            log(f"异常，跳过该学号: {e}")
            failed_list.append(student)

    with sync_playwright() as p:
        browser = launch_browser(p)
        page = open_page(browser)
        failed: List[StudentRow] = []

        log(f"批量共 {len(students)} 人，自动提交={submit}")

        for idx, student in enumerate(students, 1):
            log(f"======== [{idx}/{len(students)}] ========")
            browser, page = ensure_page(browser, page)
            # 每人独立会话，避免上一人级联框/cookie 把后面登录带到空 userId
            if idx > 1:
                close_page(page)
                page = open_page(browser)
            fill_and_collect(page, student, failed)

        if failed:
            log(f"======== 重试失败 {len(failed)} 人 ========")
            still: List[StudentRow] = []
            for i, student in enumerate(failed, 1):
                log(f"---- 重试 [{i}/{len(failed)}] ----")
                browser, page = ensure_page(browser, page)
                close_page(page)
                page = open_page(browser)
                fill_and_collect(page, student, still)
            if still:
                names = "、".join(f"{s.student_id} {s.name}".strip() for s in still)
                log(f"仍失败：{names}")
            else:
                log("重试后全部成功。")

        if not submit and keep_open_ms > 0 and page_alive(page):
            log(f"默认不提交。浏览器保留 {keep_open_ms/1000:.0f} 秒供人工检查…")
            try:
                page.wait_for_timeout(keep_open_ms)
            except Exception:
                pass

        try:
            browser.close()
        except Exception:
            pass
        log("全部结束。")


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass

    ap = argparse.ArgumentParser(description="自动填写毕业生就业质量调查问卷")
    ap.add_argument(
        "--excel",
        default=DEFAULT_EXCEL,
        help="含学号、是否QS100以内 的 Excel；存在则批量填写",
    )
    ap.add_argument(
        "--student-id",
        default="",
        help="仅填单个学号（忽略 Excel）；可配合 --qs100",
    )
    ap.add_argument("--qs100", default="否", help="单学号模式下重点高校题：是/否")
    ap.add_argument("--headless", action="store_true")
    ap.add_argument("--submit", action="store_true", help="答完后自动提交")
    ap.add_argument("--slow-mo", type=int, default=50)
    ap.add_argument("--max-steps", type=int, default=220)
    ap.add_argument("--keep-open-ms", type=int, default=5000, help="全部结束后浏览器保留毫秒数")
    ap.add_argument("--log-file", default="", help="UTF-8 日志路径")
    args = ap.parse_args()

    students: List[StudentRow] = []
    if args.student_id:
        students = [
            StudentRow(
                student_id=args.student_id.strip(),
                qs100=normalize_yes_no(args.qs100),
            )
        ]
    else:
        excel_path = Path(args.excel)
        if excel_path.is_file():
            students = load_students_from_excel(str(excel_path))
            log(f"从 Excel 读取 {len(students)} 人: {excel_path}")
            for s in students:
                log(
                    f"  {s.student_id} {s.name} 去向={s.destination or '-'} "
                    f"状态={employment_status_of(s)} 创业类型={venture_type_of(s) or '-'} "
                    f"地址={s.address or '-'} 规模={s.scale or '-'} 属于={s.belong or '-'}"
                )
        else:
            log(f"未找到 Excel: {excel_path}，改用默认学号 {DEFAULT_STUDENT_ID}")
            students = [
                StudentRow(student_id=DEFAULT_STUDENT_ID, qs100=normalize_yes_no(args.qs100))
            ]

    try:
        run(
            students,
            args.headless,
            args.submit,
            args.slow_mo,
            args.max_steps,
            args.keep_open_ms,
            args.log_file,
        )
    except PlaywrightTimeoutError as e:
        log(f"超时: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
