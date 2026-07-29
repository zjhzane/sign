import os
import sys
import time
import random
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook

# ----------------------------------------------------
# --- 【配置信息】 ---
# ----------------------------------------------------
LOGIN_URL = "https://nxg.zwu.edu.cn/api/cas/testRedirect?code=kuucNsp2rj0608fMluOtv0kwArMbaSUa4qToaiKGWt%252B%252FxBwMiP0vWpvNd6A5acLiJyuimV5QoA2KY2r79DeanVAG1AM%252FZRnWF7Um4RW30h46SWkt8kV9HUEHAcz8vPxlz77kREfBEDWPBoFUzOXLtgeolobVFX8MB12SdcvOx1Q%253D&type=pc"
ICON_SELECTOR = "#app > div.thePage > div:nth-child(7) > div.row.flex_up.flex_w > div:nth-child(5) > div.svgicon_bg.flex_center > svg > use"

FILL_SUMMARY_SELECTOR = 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[5]/div[2]/div[1]/textarea'
FINAL_RADIO_SELECTOR = 'xpath=//*[@id="app"]/div/div[2]/div[2]/form/div[2]/div[3]/div[2]/div/div/div/div[1]/span'

SUBMIT_BUTTON_SELECTOR = 'button.van-button--primary:has-text("提交")'

TIME_PICKER_YEAR = "2025年"
TIME_PICKER_CONFIRM = "button.van-picker__confirm"

LOCATION_MAPPING = {
    '办公室': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[2]/div[2]/div/div/div/div[3]/span',
    '寝室':   'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[2]/div[2]/div/div/div/div[2]/span',
    '教室':   'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[2]/div[2]/div/div/div/div[1]/span',
    '线上/电话': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[2]/div[2]/div/div/div/div[4]/span',
    '其他':   'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[2]/div[2]/div/div/div/div[5]/span',
}

THEME_MAPPING = {
    '思想引领': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[1]/span',
    '校园安全': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[2]/span',
    '学业指导': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[3]/span',
    '生涯规划': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[4]/span',
    '心理健康': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[5]/span',
    '奖励助贷': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[6]/span',
    '宗教信仰': 'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[7]/span',
    '其他':   'xpath=/html/body/div[1]/div/div[2]/div[2]/form/div[2]/div[4]/div[2]/div/div/div/div[8]/span',
}

EXCEL_FILENAME = "谈话记录模版.xlsx"
EXCEL_FILEPATH = os.path.join(os.path.expanduser('~'), 'Desktop', EXCEL_FILENAME)


# ----------------------------------------------------
# --- 【工具函数】 ---
# ----------------------------------------------------

def read_student_data_from_excel(filepath):
    if not os.path.exists(filepath):
        print(f"!!! 错误: 文件未找到: {filepath}")
        return []
    try:
        workbook = load_workbook(filepath)
        sheet = workbook.active
        data = []
        for row in range(3, sheet.max_row + 1):
            sid     = sheet[f'B{row}'].value
            cmd     = sheet[f'C{row}'].value
            theme   = sheet[f'D{row}'].value
            summary = sheet[f'E{row}'].value
            loc     = sheet[f'F{row}'].value
            if sid:
                themes = [
                    t.strip()
                    for t in str(theme or '').replace('，', ',').replace('；', ';').split(',')
                    if t.strip()
                ]
                data.append((
                    str(sid).strip(),
                    str(cmd or '').strip(),
                    themes,
                    str(summary or '').strip(),
                    str(loc or '').strip(),
                ))
        return data
    except Exception as e:
        print(f"读取Excel出错: {e}")
        return []


def click_submit(page):
    """
    找到页面真实的滚动容器滚到底部，再点击提交按钮。
    iPhone 模拟器下滚动容器是内部 div，不是 window。
    """
    # 先等提交按钮出现在页面里
    submit_btn = page.locator(SUBMIT_BUTTON_SELECTOR).last
    submit_btn.wait_for(state="attached", timeout=5000)
    # 找到页面中可滚动的容器并滚到底部
    page.evaluate("""
        () => {
            // 优先找 overflow 为 auto/scroll 的容器
            const scrollable = Array.from(document.querySelectorAll('*')).find(el => {
                const style = window.getComputedStyle(el);
                return (style.overflow === 'auto' || style.overflow === 'scroll' ||
                        style.overflowY === 'auto' || style.overflowY === 'scroll') &&
                        el.scrollHeight > el.clientHeight;
            });
            if (scrollable) {
                scrollable.scrollTop = scrollable.scrollHeight;
            } else {
                // 兜底：直接滚 document.documentElement
                document.documentElement.scrollTop = document.documentElement.scrollHeight;
                document.body.scrollTop = document.body.scrollHeight;
            }
        }
    """)
    time.sleep(0.3)
    submit_btn = page.locator(SUBMIT_BUTTON_SELECTOR).last
    submit_btn.click(force=True, timeout=3000)


# ----------------------------------------------------
# --- 【自动化主流程】 ---
# ----------------------------------------------------

def automate_workflow():
    student_list = read_student_data_from_excel(EXCEL_FILEPATH)
    if not student_list:
        return

    slow_ms = int(os.getenv("SLOW_MO", "0") or "0")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, slow_mo=slow_ms)
        context = browser.new_context(**p.devices["iPhone 12"])
        page = context.new_page()

        print("--- 正在打开页面 ---")
        page.goto(LOGIN_URL)

        try:
            page.locator(ICON_SELECTOR).click(timeout=15000)
            print("--- 已进入谈心谈话主页 ---")
        except Exception as e:
            print(f"无法点击主图标: {e}")
            return

        for i, (student_id, _time_cmd, themes, summary, loc_key) in enumerate(student_list):
            print(f"\n[学生 {i + 1}/{len(student_list)}] 学号: {student_id}")

            try:
                # 1. 进入登记页面
                print("-> 正在进入登记表单页面...")
                with page.expect_navigation(timeout=15000):
                    page.get_by_label("谈心谈话登记").get_by_text("一对一谈话").click()

                page.get_by_label("学生").wait_for(state="visible", timeout=10000)
                page.get_by_label("学生").click(force=True)

                # 2. 搜索并选择学生
                page.fill("input[placeholder='请输入学生姓名或学号搜索']", student_id)
                search_popup = page.locator(".van-popup--bottom").filter(has_text=student_id).last
                search_popup.wait_for(state="visible", timeout=5000)
                target_item = search_popup.locator(".student-search-title").filter(has_text=student_id).first
                target_item.wait_for(state="visible", timeout=5000)
                target_item.click()
                search_popup.wait_for(state="hidden", timeout=3000)
                time.sleep(0.2)

                # 3. 填充摘要
                page.fill(FILL_SUMMARY_SELECTOR, summary)

                # 4. 地点选择
                loc_selector = LOCATION_MAPPING.get(loc_key)
                if loc_selector:
                    try:
                        page.locator(loc_selector).click(timeout=3000)
                    except Exception:
                        pass

                # 5. 主题选择（多选）
                for theme in themes:
                    theme_selector = THEME_MAPPING.get(theme)
                    if theme_selector:
                        try:
                            page.locator(theme_selector).click(timeout=3000)
                        except Exception:
                            pass

                # 6. 最终单选
                try:
                    page.locator(FINAL_RADIO_SELECTOR).click(timeout=3000)
                except Exception:
                    pass

                # 7. 谈话时间
                try:
                    page.locator("input[name='talkTime']").click(timeout=5000)
                except Exception:
                    try:
                        page.get_by_label("谈话时间").click(timeout=3000)
                    except Exception:
                        pass

                # 选年份
                try:
                    page.locator("div.van-ellipsis").filter(has_text=TIME_PICKER_YEAR).first.click(timeout=5000)
                except Exception:
                    pass

                # 随机选月份和日期
                try:
                    month_labels = ["09月", "10月", "11月", "12月"]
                    picked_month = random.choice(month_labels)
                    page.locator("div.van-ellipsis").filter(has_text=picked_month).first.click(timeout=5000)
                    time.sleep(0.1)
                    day_num = random.randint(1, 30)
                    day_label = f"{day_num:02d}日"
                    page.locator("div.van-ellipsis").filter(has_text=day_label).first.click(timeout=5000)
                except Exception:
                    pass

                # 确认时间，等待弹层完全关闭
                try:
                    page.locator(TIME_PICKER_CONFIRM).click(timeout=5000)
                    page.locator(".van-overlay").last.wait_for(state="hidden", timeout=4000)
                except Exception:
                    pass

                # 8. 提交（找真实滚动容器滚底 + force 点击）
                print(f"-> 准备提交数据...")
                click_submit(page)

                # 等待跳转回列表页
                print("-> 等待跳转回列表页...")
                page.wait_for_url("**/talk-list**", timeout=20000)
                page.get_by_label("谈心谈话登记").get_by_text("一对一谈话").wait_for(state="visible", timeout=10000)

                # 等页面网络请求全部结束，确保完全稳定再开始下一个学生
                page.wait_for_load_state("networkidle", timeout=5000)

                print(f"-> [成功] 学生 {student_id} 处理完毕")

            except Exception as e:
                print(f"!!! 处理学生 {student_id} 时出错: {e}")
                page.goto(LOGIN_URL)
                page.locator(ICON_SELECTOR).click()
                time.sleep(1)

        print("\n--- 所有学生处理完毕 ---")
        input("按 Enter 键关闭浏览器...")
        browser.close()


if __name__ == "__main__":
    automate_workflow()
